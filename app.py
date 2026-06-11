import os
import re
import zipfile
import base64
import pandas as pd
from datetime import datetime
from io import BytesIO
import urllib.request
from docx import Document
from docx.shared import Pt

import streamlit as st

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT, TA_RIGHT

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", 
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

# URL del logotipo del Gobierno de Canarias para el documento PDF
URL_LOGO_CANARIAS = "https://www3.gobiernodecanarias.org/medusa/mediateca/perfeccionamiento/wp-content/uploads/sites/5/2026/06/logo-consejeria-educacion.png"

# Ocultar menús nativos de Streamlit y ajustar márgenes verticales mínimos para Moodle
st.markdown("""
    <style>
        #MainMenu {visibility: hidden;}
        footer {visibility: hidden;}
        header {visibility: hidden;}
        .block-container {padding-top: 0.1rem !important; padding-bottom: 0.1rem !important; margin-left: 0 !important; text-align: left !important;}
        div[data-testid="stVerticalBlock"] {gap: 0.3rem !important;}
        
        /* Ajustar espaciados de los componentes de entrada */
        [data-testid="stFileUploader"] {text-align: left !important; margin-bottom: 0px !important;}
        .stButton > button {width: 100% !important; margin-top: 24px !important;}
    </style>
""", unsafe_allow_html=True)

# LÓGICA DE DETECCIÓN DEL LOGO LOCAL EN GITHUB
logo_path = "logo-app.png"
if os.path.exists(logo_path):
    with open(logo_path, "rb") as image_file:
        encoded_string = base64.b64encode(image_file.read()).decode()
    img_src = f"data:image/png;base64,{encoded_string}"
else:
    img_src = "https://cdn-icons-png.flaticon.com/512/2641/2641333.png"

# CABECERA WEB COMPACTA
st.markdown(
    f"""
    <div style="font-family:'Segoe UI', Arial, sans-serif; margin-bottom: 8px; text-align: left;">
        <table style="border:none; border-collapse:collapse; width:100%; background:transparent; margin:0;">
            <tr style="border:none;">
                <td style="width:48px; vertical-align:middle; padding:0; border:none; text-align:left;">
                    <img src="{img_src}" width="42" style="display:inline-block; vertical-align:middle; border-radius:50%;"/>
                </td>
                <td style="vertical-align:middle; padding-left:10px; border:none; text-align:left;">
                    <h2 style="margin:0; color:#0A3A60; font-size:18px; font-weight:600; line-height:1.2; display:inline-block; vertical-align:middle;">
                        Consola de Certificación Oficial
                    </h2>
                    <span style="margin-left:6px; color:#718096; font-size:11.5px; display:inline-block; vertical-align:middle;">
                        &middot; Área de Formación y Perfeccionamiento
                    </span>
                </td>
            </tr>
        </table>
    </div>
    <p style='font-family:sans-serif; font-size:11.5px; color:#4A5568; margin: 0 0 10px 0; text-align:left;'>
        Soporta subida directa de archivos Excel o procesamiento síncrono mediante enlaces públicos de Google Drive.
    </p>
    """, 
    unsafe_allow_html=True
)

def limpiar_nombre_archivo(texto):
    if not texto: return ""
    return re.sub(r'[\\/*?:"<>|]', '_', str(texto).replace('\n', '').replace('\r', '').strip())

def dibujar_encabezado_pdf(canvas, doc, bytes_logo, texto_global):
    canvas.saveState()
    styles = getSampleStyleSheet()
    estilo_perf = ParagraphStyle('PerfDer', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, alignment=TA_RIGHT)
    estilo_acta = ParagraphStyle('ActaDer', parent=styles['Normal'], fontName='Helvetica', fontSize=14, alignment=TA_CENTER)
    tabla_acta_caja = Table([[Paragraph("<b>ACTA DE CERTIFICACIÓN</b>", estilo_acta)]], colWidths=[240], rowHeights=[32])
    tabla_acta_caja.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('BOX', (0,0), (-1,-1), 0.75, colors.grey),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0), ('TOPPADDING', (0,0), (-1,-1), 0),
    ]))
    contenido_derecha = [
        Paragraph("<b>PERF-06</b>", estilo_perf),
        Table([[tabla_acta_caja]], colWidths=[240], style=[('ALIGN', (0,0), (-1,-1), 'RIGHT'), ('BOTTOMPADDING', (0,0), (-1,-1), 0), ('TOPPADDING', (0,0), (-1,-1), 2)])
    ]
    if bytes_logo:
        from reportlab.lib.utils import ImageReader
        try:
            canvas.drawImage(ImageReader(BytesIO(bytes_logo)), 42, 715, width=250, height=48, preserveAspectRatio=True, mask='auto')
            celda_izquierda = Paragraph("", styles['Normal'])
        except:
            celda_izquierda = Paragraph("GOBIERNO DE CANARIAS<br/>Consejería de Educación", ParagraphStyle('Fb', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9))
    else:
        celda_izquierda = Paragraph("GOBIERNO DE CANARIAS<br/>Consejería de Educación", ParagraphStyle('Fb', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9))
    
    tabla_grafica_superior = Table([[celda_izquierda, contenido_derecha]], colWidths=[264, 264])
    tabla_grafica_superior.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'BOTTOM'), ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0), ('BOTTOMPADDING', (0,0), (-1,-1), 0), ('TOPPADDING', (0,0), (-1,-1), 0),
    ]))
    tabla_grafica_superior.wrapOn(canvas, 528, 55)
    tabla_grafica_superior.drawOn(canvas, 42, 712)
    p_intro = Paragraph(texto_global, ParagraphStyle('IntroJust', parent=styles['Normal'], fontSize=9.5, leading=14.5, alignment=TA_JUSTIFY))
    p_intro.wrapOn(canvas, 528, 120)
    p_intro.drawOn(canvas, 42, 595)
    canvas.setStrokeColor(colors.grey)
    canvas.setLineWidth(0.5)
    canvas.line(42, 582, 570, 582)
    canvas.restoreState()

def generar_acta_pdf(datos_ficha, df_coord, df_part, bytes_logo):
    buffer = BytesIO()
    f_final = datos_ficha['fecha_final']
    fecha_final_txt = f"{f_final.day} de {MESES[f_final.month - 1]} de {f_final.year}" if isinstance(f_final, datetime) else str(f_final).strip()
    f_resol = datos_ficha['fecha_resol']
    fecha_resol_txt = f_resol.strftime("%d/%m/%Y") if isinstance(f_resol, datetime) else str(f_resol).strip()
    texto_global = (
        f"Siendo las 23:59 horas del día <b>{fecha_final_txt}</b>, se da por finalizada la actividad de "
        f"Perfeccionamiento del Profesorado <b>{str(datos_ficha['nombre']).strip()}</b> realizado durante el curso escolar "
        f"<b>{str(datos_ficha['curso_escolar']).strip()}</b> en los centros educativos participantes y con nº de expediente "
        f"<b>{str(datos_ficha['exp']).strip()}</b> con un total de <b>{datos_ficha['horas_coord']}</b> horas como persona "
        f"coordinadora y <b>{datos_ficha['horas_partic']}</b> horas como participante, convocado según resolución "
        f"nº <b>{str(datos_ficha['resol']).strip()}</b> de <b>{fecha_resol_txt}</b>, en la que ha participado el profesorado "
        f"que a continuación se relaciona. En consecuencia, se propone que se expida una certificación de su "
        f"asistencia a las personas que se indican, por cumplir los requisitos que determinan la Resolución de "
        f"la Dirección General de Ordenación, Innovación y Promoción Educativa del 15 de mayo de 1998 (BOC de 8 de junio) "
        f"y la Circular de 20 de mayo de 1998."
    )
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=42, rightMargin=42, topMargin=220, bottomMargin=42)
    story = []
    styles = getSampleStyleSheet()
    estilo_cab = ParagraphStyle('TCab', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8, alignment=TA_CENTER)
    estilo_cen = ParagraphStyle('TCen', parent=styles['Normal'], fontSize=8.5, alignment=TA_CENTER)
    estilo_izq = ParagraphStyle('TIzq', parent=styles['Normal'], fontSize=8.5, alignment=TA_LEFT)
    tabla_datos = [[Paragraph("Nº", estilo_cab), Paragraph("Apellidos y Nombre", estilo_cab), Paragraph("DNI/NIF", estilo_cab), Paragraph("Rol", estilo_cab), Paragraph("Horas", estilo_cab), Paragraph("Certifica", estilo_cab)]]
    cont = 1
    for df, rol in [(df_coord, "DOCENTE COORDINADOR/A"), (df_part, "DOCENTE PARTICIPANTE")]:
        for _, fila in df.iterrows():
            tabla_datos.append([
                Paragraph(str(cont), estilo_cen), Paragraph(f"{fila.iloc[1]} {fila.iloc[2]}".strip().upper(), estilo_izq),
                Paragraph(str(fila.iloc[0]).upper(), estilo_cen), Paragraph(rol, estilo_cen),
                Paragraph(str(fila.iloc[4]), estilo_cen), Paragraph(str(fila.iloc[5]).upper(), estilo_cen)
            ])
            cont += 1
    t = Table(tabla_datos, colWidths=[40, 185, 75, 133, 40, 55], repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.lightgrey), ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('BOTTOMPADDING', (0,0), (-1,-1), 4), ('TOPPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t)
    doc.build(story, onFirstPage=lambda c, d: dibujar_encabezado_pdf(c, d, bytes_logo, texto_global),
                     onLaterPages=lambda c, d: dibujar_encabezado_pdf(c, d, bytes_logo, texto_global))
    buffer.seek(0)
    return buffer.getvalue()

def generar_memoria_oficial(datos_ficha, df_coord, df_part, bytes_plantilla):
    buffer = BytesIO()
    doc = Document(BytesIO(bytes_plantilla))
    h_coord = str(datos_ficha['horas_coord']).upper().replace("HORAS", "").strip()
    h_part = str(datos_ficha['horas_partic']).upper().replace("HORAS", "").strip()
    MAPA_REEMPLAZOS = {
        "{nombre}": str(datos_ficha['nombre']).strip(), "{exp}": str(datos_ficha['exp']).strip(),
        "{curso_escolar}": str(datos_ficha['curso_escolar']).strip(), "{horascoordinacion}": h_coord, "{horasparticipacion}": h_part
    }
    for p in doc.paragraphs:
        for cl, val in MAPA_REEMPLAZOS.items():
            if cl in p.text: p.text = p.text.replace(cl, str(val))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for cl, val in MAPA_REEMPLAZOS.items():
                    if cl in cell.text: cell.text = cell.text.replace(cl, str(val))
    tabla_certificacion = None
    for t in doc.tables:
        if len(t.rows) > 0 and any("APELLIDOS" in cell.text.upper() or "CAUSAS" in cell.text.upper() for cell in t.rows[0].cells):
            tabla_certificacion = t
            break
    if tabla_certificacion:
        cont = 1
        for df in [df_coord, df_part]:
            for _, fila in df.iterrows():
                if str(fila.iloc[5]).strip().upper() == "NO":
                    nueva_fila = tabla_certificacion.add_row()
                    nueva_fila.cells[0].text = str(cont)
                    nueva_fila.cells[1].text = f"{fila.iloc[1]}, {fila.iloc[2]}".upper()
                    nueva_fila.cells[2].text = str(fila.iloc[0]).upper()
                    nueva_fila.cells[3].text = ", ".join([str(fila.iloc[6]).strip() if pd.notna(fila.iloc[6]) else "", str(fila.iloc[7]).strip() if pd.notna(fila.iloc[7]) else ""]).strip(", ").upper()
                    for cell in nueva_fila.cells:
                        for p in cell.paragraphs:
                            for run in p.runs: run.font.size = Pt(8.5)
                    cont += 1
        if cont == 1:
            nueva_fila = tabla_certificacion.add_row()
            nueva_fila.cells[1].text = "No constan personas sin certificar"
            for p in nueva_fila.cells[1].paragraphs:
                for run in p.runs: run.font.size = Pt(8.5)
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# DISEÑO DEL PANEL DE CARGA DOBLE (HORIZONTAL OPTIMIZADO)
col_f1, col_f2, col_f3 = st.columns([1.5, 1.5, 0.8])

with col_f1:
    archivo_excel = st.file_uploader("Subir desde equipo (.xlsx)", type=["xlsx", "xls"], label_visibility="visible")

with col_f2:
    url_drive = st.text_input("O pegar enlace de Google Drive:", placeholder="https://docs.google.com/spreadsheets/d/...", label_visibility="visible")

# Determinar si hay alguna fuente de datos activa para habilitar el botón ejecutivo
fuente_activa = (archivo_excel is not None) or (url_drive.strip() != "")

with col_f3:
    ejecutar = st.button("⚡ Confeccionar", type="primary", disabled=not fuente_activa)

if ejecutar:
    with st.spinner("Procesando origen de datos..."):
        try:
            # Control primario de la plantilla en el repositorio
            if not os.path.exists("plantilla_memoria.docx"):
                st.error("Error crítico: No se encuentra 'plantilla_memoria.docx' en el repositorio de GitHub.")
                st.stop()
                
            with open("plantilla_memoria.docx", "rb") as f:
                plantilla_bytes = f.read()
            
            # Lógica de extracción de flujo si viene de Google Drive
            objeto_excel_final = None
            if archivo_excel is not None:
                objeto_excel_final = archivo_excel
            elif url_drive.strip() != "":
                match = re.search(r"/d/([a-zA-Z0-9-_]+)", url_drive)
                if match:
                    drive_id = match.group(1)
                    url_descarga_directa = f"https://docs.google.com/spreadsheets/d/{drive_id}/export?format=xlsx"
                    
                    cabeceras = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
                    req = urllib.request.Request(url_descarga_directa, headers=cabeceras)
                    objeto_excel_final = BytesIO(urllib.request.urlopen(req, timeout=8).read())
                else:
                    st.error("Error: La URL de Google Drive no contiene un ID de documento válido. Asegúrate de copiar el enlace completo de 'Compartir'.")
                    st.stop()
            
            if objeto_excel_final is None:
                st.error("Por favor, suba un archivo o provea un enlace válido.")
                st.stop()

            # Intento de descarga del logo institucional para el PDF
            bytes_logo_pdf = None
            try:
                cabeceras_logo = {'User-Agent': 'Mozilla/5.0'}
                req_logo = urllib.request.Request(URL_LOGO_CANARIAS, headers=cabeceras_logo)
                bytes_logo_pdf = urllib.request.urlopen(req_logo, timeout=4).read()
            except:
                pass
            
            # Carga de hojas del libro de cálculo
            df_ficha = pd.read_excel(objeto_excel_final, sheet_name="Ficha del Proyecto", header=None)
            df_coord = pd.read_excel(objeto_excel_final, sheet_name="Coordinador")
            df_part = pd.read_excel(objeto_excel_final, sheet_name="Participante")
            
            datos_ficha = {
                "nombre": df_ficha.iloc[2, 2], "exp": str(df_ficha.iloc[4, 2]),
                "resol": str(df_ficha.iloc[4, 6]), "fecha_resol": df_ficha.iloc[5, 6],
                "horas_coord": df_ficha.iloc[7, 2], "horas_partic": df_ficha.iloc[8, 2],
                "curso_escolar": df_ficha.iloc[9, 6], "fecha_final": df_ficha.iloc[11, 2]
            }
            
            pdf_bytes = generar_acta_pdf(datos_ficha, df_coord, df_part, bytes_logo_pdf)
            docx_bytes = generar_memoria_oficial(datos_ficha, df_coord, df_part, plantilla_bytes)
            
            zip_buffer = BytesIO()
            exp_limpio = limpiar_nombre_archivo(datos_ficha['exp'])
            nom_limpio = limpiar_nombre_archivo(datos_ficha['nombre'])
            
            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                zip_file.writestr(f"Perf06_{exp_limpio}_{nom_limpio}.pdf", pdf_bytes)
                zip_file.writestr(f"Memoria_{exp_limpio}.docx", docx_bytes)
            zip_buffer.seek(0)
            
            col_res1, col_res2 = st.columns([2, 1])
            with col_res1:
                st.success("✨ ¡Paquete generado con éxito!")
            with col_res2:
                st.download_button(
                    label="📥 Descargar Paquete (.ZIP)",
                    data=zip_buffer.getvalue(),
                    file_name=f"Certificacion_Proyecto_{exp_limpio}.zip",
                    mime="application/zip"
                )
                
        except Exception as e:
            st.error(f"Error crítico de procesamiento: {str(e)}")
            st.info("Nota de Drive: Si usó un enlace de Google Drive, asegúrese de que esté configurado como 'Cualquier persona con el enlace puede ver'.")

# Pie de página y RGPD
st.markdown(
    """
    <div style="margin-top: 15px; font-family: sans-serif; font-size: 10px; color: #A0AEC0; text-align: left; line-height: 1.2;">
        🔒 <b>RGPD:</b> Datos procesados en memoria RAM volátil y destruidos al finalizar de forma inmediata.
        <br/><span style="font-weight: bold; font-size: 8px; letter-spacing: 0.5px;">DEVELOPER 1.1 (DRIVE INTEGRATION)</span>
    </div>
    """, 
    unsafe_allow_html=True
)